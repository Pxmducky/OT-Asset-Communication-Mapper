from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.services.asset_service import AssetService
from app.services.communication_service import CommunicationService

ASSET_COLUMNS = [
    "asset_code", "asset_name", "asset_type", "ip", "mac", "hostname", "vlan",
    "vendor", "product", "plant", "building", "production_line", "description",
]

COMMUNICATION_COLUMNS = [
    "source_asset", "destination_asset", "source", "destination", "protocol",
    "source_port", "destination_port", "source_name", "destination_name", "description",
]


class ExcelExporter:
    """Genera el inventario completo en Excel.

    Usa valores (no fórmulas), así el mismo archivo se puede volver a importar.
    """

    def __init__(self, db: Session):
        self.db = db

    def export(self) -> BytesIO:
        workbook = Workbook()

        assets_sheet = workbook.active
        assets_sheet.title = "Assets"
        assets = AssetService.get_all(self.db)
        self._write_sheet(
            assets_sheet,
            ASSET_COLUMNS,
            [[getattr(asset, column) for column in ASSET_COLUMNS] for asset in assets],
        )

        communications_sheet = workbook.create_sheet("Communications")
        rows = [
            [
                communication.source_asset.asset_code,
                communication.destination_asset.asset_code,
                communication.source_ip,
                communication.destination_ip,
                communication.protocol,
                communication.source_port,
                communication.destination_port,
                communication.source_name,
                communication.destination_name,
                communication.description,
            ]
            for communication in CommunicationService.get_all(self.db)
        ]
        self._write_sheet(communications_sheet, COMMUNICATION_COLUMNS, rows)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def _write_sheet(sheet, headers: list[str], rows: list[list]) -> None:
        sheet.append(headers)
        for row in rows:
            sheet.append(row)

        header_font = Font(name="Arial", bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", start_color="1F2937")
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for index, header in enumerate(headers, start=1):
            longest = max([len(header)] + [len(str(row[index - 1] or "")) for row in rows[:1000]])
            sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, 60)